"""
Handles ModelTask.EXTRACT_TEXT_NATIVE — the "basic, deterministic text
extraction" path from the architecture: PyMuPDF for PDFs with a real text
layer, python-docx for DOCX, and Tesseract OCR as the fallback when a PDF
turns out to have no usable text layer (i.e. it's actually a scan).

This processor does NOT attempt layout/table/form understanding — that's
QwenVisionProcessor's job (ModelTask.EXTRACT_VISUAL). If a document needs
that kind of understanding, the caller should route to EXTRACT_VISUAL
directly rather than relying on this processor to escalate internally.
"""

import asyncio
import io
import logging

import pymupdf
import pytesseract
from docx import Document as DocxDocument
from PIL import Image

from app.core.config import settings
from app.services.ai_brain.context import ProcessingContext, ProcessorResult
from app.services.ai_brain.processors.base import AIProcessor
from app.services.ai_brain.tasks import ModelTask

logger = logging.getLogger(__name__)

# Below this average chars-per-page, treat a PDF as having no real text
# layer (i.e. it's a scan saved as PDF) and fall back to OCR.
MIN_CHARS_PER_PAGE_FOR_NATIVE_TEXT = 20

DOCX_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class TesseractProcessor(AIProcessor):
    task = ModelTask.EXTRACT_TEXT_NATIVE

    def __init__(self) -> None:
        if settings.TESSERACT_CMD:
            pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_CMD

    async def run(self, context: ProcessingContext) -> ProcessorResult:
        if not context.image_bytes:
            return ProcessorResult(
                success=False,
                task=self.task,
                model_name="tesseract",
                error="No file bytes provided in context.image_bytes",
            )

        try:
            # PyMuPDF/pytesseract/python-docx are all synchronous/CPU-bound —
            # run in a thread so this doesn't block the event loop.
            text, confidence, used_ocr = await asyncio.to_thread(
                self._extract, context.image_bytes, context.mime_type
            )
        except Exception as exc:  # noqa: BLE001 — deliberately broad: any
            # extraction failure should become a NEEDS_REVIEW-routable
            # result, not an unhandled exception bubbling out of the router.
            logger.exception("Text extraction failed")
            return ProcessorResult(
                success=False,
                task=self.task,
                model_name="tesseract",
                error=str(exc),
            )

        return ProcessorResult(
            success=True,
            task=self.task,
            model_name="tesseract" if used_ocr else "pymupdf",
            output={"text": text, "used_ocr": used_ocr},
            confidence=confidence,
        )

    def _extract(
        self, file_bytes: bytes, mime_type: str | None
    ) -> tuple[str, float, bool]:
        if mime_type == "application/pdf":
            return self._extract_pdf(file_bytes)

        if mime_type in DOCX_MIME_TYPES:
            return self._extract_docx(file_bytes)

        if mime_type and mime_type.startswith("image/"):
            return self._ocr_image(file_bytes)

        raise ValueError(
            f"Unsupported mime_type for text extraction: {mime_type}"
        )

    def _extract_pdf(self, pdf_bytes: bytes) -> tuple[str, float, bool]:
        doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")

        try:
            pages_text = [page.get_text() for page in doc]

            avg_chars = sum(len(t) for t in pages_text) / max(
                len(pages_text), 1
            )

            if avg_chars >= MIN_CHARS_PER_PAGE_FOR_NATIVE_TEXT:
                return (
                    "\n".join(pages_text).strip(),
                    1.0,
                    False,
                )

            # No usable text layer — rasterize each page and OCR it.
            ocr_texts: list[str] = []
            confidences: list[float] = []

            for page in doc:
                pix = page.get_pixmap(dpi=200)
                image = Image.open(
                    io.BytesIO(pix.tobytes("png"))
                )

                text, conf = self._ocr_image_object(image)

                ocr_texts.append(text)
                confidences.append(conf)

            overall_confidence = (
                sum(confidences) / len(confidences)
                if confidences
                else 0.0
            )

            return (
                "\n".join(ocr_texts).strip(),
                overall_confidence,
                True,
            )

        finally:
            doc.close()

    def _extract_docx(self, docx_bytes: bytes) -> tuple[str, float, bool]:
        doc = DocxDocument(io.BytesIO(docx_bytes))

        paragraphs = [
            p.text
            for p in doc.paragraphs
            if p.text.strip()
        ]

        table_cells = [
            cell.text
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
            if cell.text.strip()
        ]

        text = "\n".join(paragraphs + table_cells).strip()

        return text, 1.0, False

    def _ocr_image(
        self, image_bytes: bytes
    ) -> tuple[str, float, bool]:
        image = Image.open(io.BytesIO(image_bytes))

        text, confidence = self._ocr_image_object(image)

        return text, confidence, True

    def _ocr_image_object(
        self, image: Image.Image
    ) -> tuple[str, float]:
        data = pytesseract.image_to_data(
            image,
            lang=settings.OCR_LANGUAGE,
            output_type=pytesseract.Output.DICT,
        )

        words = [
            w
            for w in data["text"]
            if w.strip()
        ]

        confs = [
            float(c)
            for c, w in zip(data["conf"], data["text"])
            if w.strip() and float(c) >= 0
        ]

        text = " ".join(words)

        # Tesseract reports 0-100; normalize to 0-1
        # to match the rest of the pipeline.
        avg_confidence = (
            sum(confs) / len(confs) / 100.0
            if confs
            else 0.0
        )

        return text, avg_confidence